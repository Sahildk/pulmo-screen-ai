from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import json

doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Helper functions
def add_heading_h1(text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(23)

def add_heading_h2(text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(17)

def add_heading_h3(text):
    p = doc.add_heading(text, level=3)
    for run in p.runs:
        run.font.name = 'Times New Roman'

def add_heading_h4(text):
    p = doc.add_heading(text, level=4)
    for run in p.runs:
        run.font.name = 'Times New Roman'

def add_heading_h5(text):
    p = doc.add_heading(text, level=5)
    for run in p.runs:
        run.font.name = 'Times New Roman'

def add_normal(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    return p

def add_mixed(parts):
    """parts is a list of (text, bold) tuples"""
    p = doc.add_paragraph()
    for text, bold in parts:
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = bold
    return p

def add_bullet(text, bold=False):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    return p

def add_code_line(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.bold = True
    return p

# ===== PHASE 1 =====
add_heading_h1("Phase 1: Lung Cancer Risk Prediction using Machine Learning")

# Group Members
add_normal("Group Members:", bold=True)

# Group Members Table
table = doc.add_table(rows=3, cols=4)
table.style = 'Table Grid'
members = [
    ("01", "Nikhil Ankola", "TYIT-1", "23UF18042IT001"),
    ("02", "Shriya Bhambure", "TYIT-1", "23UF18375IT005"),
    ("07", "Sahil Deore", "TYIT-1", "23UF17901IT013"),
]
for i, (sr, name, div, roll) in enumerate(members):
    table.rows[i].cells[0].text = sr
    table.rows[i].cells[1].text = name
    table.rows[i].cells[2].text = div
    table.rows[i].cells[3].text = roll
    for cell in table.rows[i].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)

doc.add_paragraph("")

# 1. Project Title
add_heading_h2("1. Project Title")
add_normal("Lung Cancer Risk Prediction using Machine Learning (Multiple Classification Algorithms)", bold=True)
doc.add_paragraph("")

# 2. Problem Statement
add_heading_h2("2. Problem Statement")
add_mixed([
    ("Lung cancer is one of the ", False),
    ("most common and deadliest forms of cancer", True),
    (" worldwide. It accounts for millions of deaths annually and is often diagnosed at a late stage when treatment options are limited. A significant number of individuals remain ", False),
    ("undiagnosed", True),
    (" until the disease reaches an advanced, often fatal stage.", False),
])
add_mixed([
    ("Traditional diagnostic methods rely on ", False),
    ("imaging scans, biopsies, and clinical evaluation", True),
    (", which may be ", False),
    ("time-consuming, costly, and not always accessible", True),
    (". These methods often fail to provide early warning signs, creating a need for a more efficient and intelligent solution.", False),
])
add_mixed([
    ("This project aims to develop a ", False),
    ("machine learning-based predictive model", True),
    (" that can identify individuals who may be at risk of lung cancer by analyzing various ", False),
    ("health and lifestyle factors such as smoking habits, age, air pollution exposure, chest pain, breathing issues, and other symptoms", True),
    (".", False),
])
add_normal("The main objective of this system is to assist in early risk prediction, support healthcare professionals in decision-making, and promote preventive healthcare awareness among individuals.")
doc.add_paragraph("")

# 3. Introduction
add_heading_h2("3. Introduction")
add_normal("In the modern era, machine learning has emerged as a transformative technology in healthcare. It allows systems to learn from data and make accurate predictions without explicit programming. By analyzing large medical datasets, machine learning can identify patterns and provide early warnings about potential health risks.")
add_normal("Lung cancer is one of the most prevalent and dangerous diseases, influenced by factors such as smoking, air pollution exposure, alcohol consumption, and various respiratory symptoms. Early detection is essential to reduce complications and improve patient outcomes.")
add_normal("In this project, a structured health and lifestyle dataset is used for building the predictive model. The dataset contains multiple health predictor variables and a target variable (LUNG_CANCER) indicating whether a person is at high risk (YES) or low risk (NO).")
add_normal("This project uses supervised machine learning techniques, particularly classification algorithms, to analyze the dataset.")
add_normal("Multiple classification algorithms are trained and compared:")
add_bullet("Random Forest")
add_bullet("Support Vector Machine (SVM)")
add_bullet("Logistic Regression")
add_bullet("K-Nearest Neighbors (KNN)")
add_bullet("Decision Tree")
add_normal("The best performing model is selected based on F1-score, ensuring a balance between precision and recall.")
add_normal("This ensemble approach of comparing multiple algorithms ensures the most optimal model is selected, resulting in improved performance and reliability.")
add_normal("This project demonstrates how data-driven approaches can improve healthcare systems and assist in solving real-world problems effectively.")
doc.add_paragraph("")

# 4. Scope of the Project
add_heading_h2("4. Scope of the Project")
add_normal("The scope of this project includes the development of a machine learning-based system for predicting lung cancer risk using structured health and lifestyle data.")
add_normal("The system is designed to analyze patient data and classify individuals into:")
add_bullet("High Risk of Lung Cancer")
add_bullet("Low Risk of Lung Cancer")

add_heading_h3("Key Components of the Project")
add_mixed([
    ("Data Collection:\n", True),
    (" The dataset contains 309 records with 16 health and lifestyle features, already preprocessed and cleaned.", False),
])
add_mixed([
    ("Data Preprocessing:\n", True),
    (" Includes label encoding of categorical variables (Yes/No, M/F), handling class imbalance using SMOTE oversampling, and stratified train-test splitting.", False),
])
add_mixed([
    ("Feature Analysis:\n", True),
    (" Identifying important health factors such as smoking, fatigue, allergy, coughing, and swallowing difficulty that influence predictions.", False),
])
add_mixed([
    ("Model Implementation:\n", True),
    (" Applying multiple machine learning algorithms:", False),
])
add_bullet("Random Forest")
add_bullet("Support Vector Machine (SVM)")
add_bullet("Logistic Regression")
add_bullet("K-Nearest Neighbors (KNN)")
add_bullet("Decision Tree")
add_mixed([
    ("Model Evaluation:\n", True),
    (" Using performance metrics such as:", False),
])
add_bullet("Accuracy")
add_bullet("Precision")
add_bullet("Recall")
add_bullet("F1 Score")
add_bullet("Confusion Matrix")

add_heading_h3("Future Scope")
add_normal("The project can be extended in the following ways:")
add_bullet("Development of a web or mobile application for real-time prediction")
add_bullet("Use of advanced algorithms or deep learning techniques")
add_bullet("Integration with hospital or healthcare systems")
add_bullet("Use of larger and more diverse datasets for improved accuracy")

add_heading_h3("Applications")
add_bullet("Early detection of lung cancer risk")
add_bullet("Healthcare decision support systems")
add_bullet("Preventive healthcare awareness")
add_bullet("Educational and research purposes")

# 5. SDG Goal Alignment
add_heading_h2("5. SDG Goal Alignment")
add_normal("This project aligns with United Nations Sustainable Development Goal (SDG) 3: Good Health and Well-being.")
add_normal("SDG 3 aims to ensure healthy lives and promote well-being for all age groups. One of its major focuses is reducing the impact of non-communicable diseases like lung cancer through prevention and early diagnosis.")
doc.add_paragraph("")

add_heading_h3("Project Contribution to SDG 3")
add_mixed([("Early Diagnosis:\n", True), (" Helps in identifying high-risk individuals at an early stage.", False)])
add_mixed([("Preventive Healthcare:\n", True), (" Encourages individuals to take timely medical and lifestyle actions.", False)])
add_mixed([("Reduced Healthcare Costs:\n", True), (" Minimizes expensive treatments through early intervention.", False)])
add_mixed([("Decision Support:\n", True), (" Assists healthcare professionals with data-driven insights.", False)])
add_mixed([("Technological Advancement:\n", True), (" Promotes the use of machine learning in healthcare systems.", False)])

doc.add_paragraph("")

# ===== PHASE 2 =====
add_heading_h1("Phase 2:")

# 6. Design and Methodology
add_heading_h3("6. Design and Methodology")
add_heading_h4("System Architecture Overview:")
add_normal("The proposed system follows a structured machine learning pipeline based on supervised learning principles. The model is trained on labeled data (high risk or low risk outcomes) and learns patterns to predict lung cancer risk for new, unseen cases.")
add_normal("Workflow Components:")
add_bullet("Data collection from structured health dataset (Dataset.xlsx)")
add_bullet("Data preprocessing and encoding")
add_bullet("Class imbalance handling using SMOTE")
add_bullet("Feature selection and analysis")
add_bullet("Model training using multiple classifiers")
add_bullet("Model evaluation using standard metrics")
add_bullet("Best model selection based on F1-score")

add_heading_h4("Data Preprocessing:")
add_normal("Data preprocessing is critical to ensure model accuracy and efficiency. The following steps are implemented:")
add_normal("Label Encoding", bold=True)
add_normal("All categorical features (GENDER, SMOKING, YELLOW_FINGERS, ANXIETY, etc.) contain Yes/No or M/F values which are converted to numeric form using LabelEncoder. This ensures the ML algorithms can process the data effectively.")
add_normal("Handling Class Imbalance", bold=True)
add_normal("The dataset is heavily imbalanced with 270 YES cases (87%) and 39 NO cases (13%). SMOTE (Synthetic Minority Over-sampling Technique) is applied on the training set to generate synthetic samples of the minority class, ensuring the model learns both classes equally.")
add_normal("Data Partitioning", bold=True)
add_normal("The dataset is divided as follows:")
add_bullet("Training Set: 80% (model learning)")
add_bullet("Testing Set: 20% (model evaluation)")
add_normal("Stratified splitting is used to maintain the same class distribution in both training and testing sets. This ensures the model is evaluated on unseen data, providing an unbiased performance assessment.")

add_heading_h4("Feature Selection and Analysis:")
add_normal("Feature selection identifies the most influential attributes for lung cancer risk prediction, improving model interpretability and computational efficiency.")
add_normal("Selected Features:", bold=True)
features = [
    "Gender (M/F)", "Age (years)", "Smoking (Yes/No)", "Yellow Fingers (Yes/No)",
    "Anxiety (Yes/No)", "Peer Pressure (Yes/No)", "Chronic Disease (Yes/No)",
    "Fatigue (Yes/No)", "Allergy (Yes/No)", "Wheezing (Yes/No)",
    "Alcohol Consuming (Yes/No)", "Coughing (Yes/No)", "Shortness of Breath (Yes/No)",
    "Swallowing Difficulty (Yes/No)", "Chest Pain (Yes/No)", "Age Group (Young/Middle/Old)"
]
for f in features:
    add_bullet(f)

add_normal("Selection Techniques:")
add_bullet("Logistic Regression coefficient analysis")
add_bullet("Feature importance scores from normalized model coefficients")
add_normal("Key Findings: Swallowing Difficulty, Allergy, Chronic Disease, Coughing, and Yellow Fingers emerge as highly influential features with the strongest predictive power for lung cancer risk.", bold=True)

add_heading_h4("Machine Learning Algorithm Selection:")

# Multiple algorithms
for algo, desc in [
    ("Logistic Regression (Best Model)", "A linear classification algorithm that models the probability of a binary outcome. It uses a logistic (sigmoid) function to map predictions to probabilities. Selected as the best model with F1-score of 94.34%."),
    ("Random Forest", "An ensemble learning method that builds multiple decision trees and combines their predictions using majority voting, resulting in improved performance and stability."),
    ("Support Vector Machine (SVM)", "A classification algorithm that finds the optimal hyperplane to separate different classes in high-dimensional space. Uses kernel functions to handle non-linear data."),
    ("K-Nearest Neighbors (KNN)", "A simple yet effective algorithm that classifies new data points based on the majority class of their k nearest neighbors in the feature space."),
    ("Decision Tree", "A tree-structured classifier where internal nodes represent feature tests, branches represent outcomes, and leaf nodes represent class labels."),
]:
    add_heading_h5(algo)
    add_normal(desc)

doc.add_paragraph("")

# 7. Implementation and Results
add_heading_h3("7. Implementation and Results")
add_heading_h4("Implementation Environment:")
add_normal("Programming Language: Python 3.x")
add_normal("Libraries Utilized:")
add_bullet("pandas: Data manipulation and analysis")
add_bullet("numpy: Numerical computations")
add_bullet("scikit-learn: Machine learning algorithms and evaluation metrics")
add_bullet("imbalanced-learn: SMOTE oversampling for class imbalance")
add_bullet("Flask: REST API backend server")
add_bullet("React (Vite): Frontend user interface")
add_bullet("Recharts: Data visualization and charts")

add_heading_h4("Output Results:")
doc.add_paragraph("")  # Placeholder for screenshots

add_heading_h4("Performance Metrics and Evaluation:")
doc.add_paragraph("")
add_normal("The model's performance is assessed using industry-standard metrics:")

# Accuracy
add_normal("Accuracy", bold=True)
add_normal("Definition: Overall correctness of predictions")
add_normal("Formula: (TP + TN) / (TP + TN + FP + FN)")
add_normal("Best Model Result: 90.32% (Logistic Regression)")
add_normal("Interpretation: Higher accuracy indicates more reliable predictions")
doc.add_paragraph("")

# Confusion Matrix
add_normal("Confusion Matrix", bold=True)
add_normal("This matrix reveals the breakdown of correct and incorrect predictions, essential for understanding model behavior.")

# Confusion matrix table
cm_table = doc.add_table(rows=3, cols=3)
cm_table.style = 'Table Grid'
cm_table.rows[0].cells[0].text = ""
cm_table.rows[0].cells[1].text = "Predicted Negative (NO)"
cm_table.rows[0].cells[2].text = "Predicted Positive (YES)"
cm_table.rows[1].cells[0].text = "Actual Negative (NO)"
cm_table.rows[1].cells[1].text = "True Negatives (TN) = 6"
cm_table.rows[1].cells[2].text = "False Positives (FP) = 2"
cm_table.rows[2].cells[0].text = "Actual Positive (YES)"
cm_table.rows[2].cells[1].text = "False Negatives (FN) = 4"
cm_table.rows[2].cells[2].text = "True Positives (TP) = 50"
for row in cm_table.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)

doc.add_paragraph("")

# Precision
add_normal("Precision", bold=True)
add_normal("Formula: TP / (TP + FP)")
add_normal("Definition: Proportion of positive predictions that are correct")
add_normal("Result: 96.15%")
add_normal("Importance: Critical in medical contexts to minimize false lung cancer diagnoses")

# Recall
add_normal("Recall (Sensitivity)", bold=True)
add_normal("Formula: TP / (TP + FN)")
add_normal("Definition: Ability to identify actual lung cancer cases")
add_normal("Result: 92.59%")
add_normal("Importance: Ensures patients with lung cancer risk are not missed")

# F1 Score
add_normal("F1 Score", bold=True)
add_normal("Formula: 2 x (Precision x Recall) / (Precision + Recall)")
add_normal("Definition: Harmonic mean balancing precision and recall")
add_normal("Result: 94.34%")
add_normal("Usage: Provides single metric for model comparison")
doc.add_paragraph("")

# Model Comparison Table
add_heading_h4("Model Comparison Results:")
comp_table = doc.add_table(rows=6, cols=5)
comp_table.style = 'Table Grid'
headers = ["Algorithm", "Accuracy", "Precision", "Recall", "F1 Score"]
for j, h in enumerate(headers):
    comp_table.rows[0].cells[j].text = h
    for p in comp_table.rows[0].cells[j].paragraphs:
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.bold = True

models_data = [
    ("Random Forest", 87.10, 92.59, 92.59, 92.59),
    ("SVM", 45.16, 85.71, 44.44, 58.54),
    ("Logistic Regression (Best)", 90.32, 96.15, 92.59, 94.34),
    ("KNN", 83.87, 94.00, 87.04, 90.38),
    ("Decision Tree", 85.48, 90.91, 92.59, 91.74),
]
for i, (name, acc, prec, rec, f1) in enumerate(models_data):
    comp_table.rows[i+1].cells[0].text = name
    comp_table.rows[i+1].cells[1].text = f"{acc:.2f}%"
    comp_table.rows[i+1].cells[2].text = f"{prec:.2f}%"
    comp_table.rows[i+1].cells[3].text = f"{rec:.2f}%"
    comp_table.rows[i+1].cells[4].text = f"{f1:.2f}%"
    for cell in comp_table.rows[i+1].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

doc.add_paragraph("")

# 8. Code / Tool Usage
add_heading_h4("8. Code / Tool Usage")
add_heading_h4("Execution Pipeline:")
add_normal("The implementation follows these steps:")
add_bullet("Dataset Loading: Excel file imported using pandas")
add_bullet("Data Preprocessing: Categorical encoding via LabelEncoder, column name cleaning")
add_bullet("Class Imbalance Handling: SMOTE applied to training data")
add_bullet("Data Partitioning: Training and testing sets created with stratified split")
add_bullet("Model Training: Five classifiers trained on resampled training data")
add_bullet("Prediction: Models generate predictions on test data")
add_bullet("Performance Evaluation: Metrics calculated and compared")
add_bullet("Best Model Selection: Logistic Regression selected based on highest F1-score")

add_heading_h4("Python Implementation Code:")

# Read the actual code files
with open('backend/train_model.py', 'r') as f:
    train_code = f.read()
with open('backend/app.py', 'r') as f:
    app_code = f.read()

add_normal("Training Script (train_model.py):", bold=True)
for line in train_code.split('\n'):
    add_code_line(line)

doc.add_paragraph("")
add_normal("Flask API Server (app.py):", bold=True)
for line in app_code.split('\n'):
    add_code_line(line)

doc.add_paragraph("")

# 9. Conclusion
add_heading_h3("9. Conclusion")
add_normal("This project successfully demonstrates a machine learning-based approach for lung cancer risk prediction. The Logistic Regression algorithm proves most effective among five tested classifiers, achieving an F1-score of 94.34% with 96.15% precision and 92.59% recall.")

add_normal("Key Achievements:", bold=True)
add_bullet("Robust preprocessing pipeline handling categorical encoding and class imbalance using SMOTE")
add_bullet("Effective feature analysis identifying key predictive variables (Swallowing Difficulty, Allergy, Chronic Disease, Coughing)")
add_bullet("Comprehensive comparison of five ML algorithms with the best automatically selected")
add_bullet("Full-stack web application with React frontend and Flask API backend for real-time predictions")
add_bullet("Comprehensive evaluation demonstrating clinical viability")

add_normal("Future Enhancements:", bold=True)
add_bullet("Integration of additional medical features (CT scan data, genetic markers)")
add_bullet("Cross-validation and hyperparameter optimization")
add_bullet("Deep learning approaches for improved accuracy")
add_bullet("Regular model retraining with new patient data")

add_normal("The methodology presented aligns with industry standards and can effectively support early lung cancer diagnosis and preventive healthcare initiatives.")

# Save
doc.save("MLBI_Stage_2_LungCancer.docx")
print("Document generated successfully: MLBI_Stage_2_LungCancer.docx")
